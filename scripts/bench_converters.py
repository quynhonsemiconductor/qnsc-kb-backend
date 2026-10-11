"""Measure structure yield and wall clock for document-to-Markdown converters.

Run against the fixtures from bench_make_fixtures.py. Reports, per converter:
wall clock, output characters, Markdown heading count, Markdown table rows, and
numeric-token coverage against the plain page extractor (the lossless baseline).
"""

from __future__ import annotations

import io
import re
import sys
import time
from collections import Counter
from pathlib import Path

NUMERIC = re.compile(r"(?<!\d)\d+(?:[.,]\d+)*(?:\s?[%％])?(?!\d)")


def headings(markdown: str) -> int:
    return len(re.findall(r"^#{1,6}\s+\S", markdown, re.MULTILINE))


def table_rows(markdown: str) -> int:
    return len([line for line in markdown.splitlines() if line.strip().startswith("|")])


def bullets(markdown: str) -> int:
    return len(re.findall(r"^\s*[-*+]\s+\S", markdown, re.MULTILINE))


def numeric_coverage(original: str, candidate: str) -> float:
    original_counts = Counter(m.group(0).strip() for m in NUMERIC.finditer(original))
    if not original_counts:
        return 1.0
    candidate_counts = Counter(m.group(0).strip() for m in NUMERIC.finditer(candidate))
    preserved = sum(
        min(count, candidate_counts.get(token, 0))
        for token, count in original_counts.items()
    )
    return preserved / sum(original_counts.values())


def token_coverage(original: str, candidate: str) -> float:
    pattern = r"[A-Za-zÀ-ỹ0-9][A-Za-zÀ-ỹ0-9_-]{3,}"
    original_tokens = set(re.findall(pattern, original.lower()))
    if not original_tokens:
        return 1.0
    candidate_tokens = set(re.findall(pattern, candidate.lower()))
    return len(original_tokens & candidate_tokens) / len(original_tokens)


# --- converters ---------------------------------------------------------------


def plain_pages(path: Path, data: bytes) -> str:
    """The current lossless baseline: pypdf / python-docx text, no structure."""
    if path.suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def markitdown_convert(path: Path, data: bytes) -> str:
    from markitdown import MarkItDown

    converter = MarkItDown(enable_plugins=False)
    result = converter.convert_stream(io.BytesIO(data), file_extension=path.suffix)
    return (
        getattr(result, "markdown", None) or getattr(result, "text_content", "") or ""
    )


def pymupdf4llm_convert(path: Path, data: bytes) -> str:
    if path.suffix != ".pdf":
        raise RuntimeError("pdf only")
    import pymupdf
    import pymupdf4llm

    document = pymupdf.open(stream=data, filetype="pdf")
    return pymupdf4llm.to_markdown(document, show_progress=False)


CONVERTERS = {
    "page_extractor": plain_pages,
    "markitdown": markitdown_convert,
    "pymupdf4llm": pymupdf4llm_convert,
}


def main() -> None:
    paths = [Path(p) for p in sys.argv[1:]]
    for path in paths:
        data = path.read_bytes()
        baseline = plain_pages(path, data)
        print(f"\n=== {path.name}  ({len(data):,} bytes) ===")
        print(
            f"{'converter':16} {'secs':>7} {'chars':>8} {'head':>5} {'tbl':>5} "
            f"{'bul':>5} {'tok%':>6} {'num%':>6}"
        )
        for name, fn in CONVERTERS.items():
            try:
                start = time.perf_counter()
                output = fn(path, data)
                elapsed = time.perf_counter() - start
            except Exception as exc:
                print(f"{name:16} {'-':>7}  {type(exc).__name__}: {exc}")
                continue
            print(
                f"{name:16} {elapsed:7.2f} {len(output):8,} {headings(output):5} "
                f"{table_rows(output):5} {bullets(output):5} "
                f"{token_coverage(baseline, output) * 100:6.1f} "
                f"{numeric_coverage(baseline, output) * 100:6.1f}"
            )


if __name__ == "__main__":
    main()
