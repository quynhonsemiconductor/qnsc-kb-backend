"""Build structured test documents for converter benchmarking.

Writes a PDF (headings, paragraphs, a real table) and a DOCX with the same
content so converter output can be compared on identical material.
"""

from __future__ import annotations

import sys
from pathlib import Path

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "bench_fixtures")
OUT.mkdir(parents=True, exist_ok=True)

TITLE = "Quy trinh kiem tra chat luong wafer"
SECTIONS = [
    (
        "1. Pham vi ap dung",
        [
            "Quy trinh nay ap dung cho toan bo cong doan kiem tra wafer tai nha may QNSC.",
            "Moi lo hang phai duoc kiem tra truoc khi chuyen sang cong doan dong goi.",
            "Nguoi chiu trach nhiem: Truong bo phan QA, ma tai lieu QA-WF-001, phien ban 2.3.",
        ],
    ),
    (
        "2. Thiet bi do luong",
        [
            "May do do day su dung dau do 4 diem, sai so toi da 0.5%.",
            "Kinh hien vi quang hoc do phong dai 1000x de kiem tra khuyet tat be mat.",
            "Thiet bi phai duoc hieu chuan moi 90 ngay theo tieu chuan ISO 9001:2015.",
        ],
    ),
    (
        "3. Nguong chap nhan",
        [
            "Lo hang bi tu choi khi ty le khuyet tat vuot qua 2.5% tong so wafer.",
            "Do day trung binh phai nam trong khoang 725 um +/- 15 um.",
            "Dien tro suat yeu cau tu 8.5 den 11.5 ohm-cm.",
        ],
    ),
]
TABLE = [
    ["Thong so", "Gia tri toi thieu", "Gia tri toi da", "Don vi"],
    ["Do day", "710", "740", "um"],
    ["Dien tro suat", "8.5", "11.5", "ohm-cm"],
    ["Ty le khuyet tat", "0", "2.5", "%"],
    ["Do nham be mat", "0.1", "0.8", "nm"],
]


def build_pdf(path: Path, repeat: int) -> None:
    import pymupdf

    document = pymupdf.open()
    for copy in range(repeat):
        page = document.new_page()
        writer = pymupdf.TextWriter(page.rect)
        bold = pymupdf.Font("hebo")
        plain = pymupdf.Font("helv")
        y = 60.0
        if copy == 0:
            writer.append((56, y), TITLE, font=bold, fontsize=19)
            y += 34
        for heading, lines in SECTIONS:
            writer.append((56, y), heading, font=bold, fontsize=14)
            y += 22
            for line in lines:
                writer.append((56, y), line, font=plain, fontsize=10.5)
                y += 16
            y += 8
        writer.append((56, y), "4. Bang thong so ky thuat", font=bold, fontsize=14)
        y += 24
        # Draw a real ruled table so table-aware converters can detect it.
        col_x = [56, 196, 316, 436, 520]
        row_h = 20
        for row_index, row in enumerate(TABLE):
            for col_index, cell in enumerate(row):
                font = bold if row_index == 0 else plain
                writer.append(
                    (col_x[col_index] + 4, y + 14), cell, font=font, fontsize=9.5
                )
            y += row_h
        table_top = y - row_h * len(TABLE)
        writer.write_text(page)
        for row_index in range(len(TABLE) + 1):
            line_y = table_top + row_index * row_h
            page.draw_line(
                pymupdf.Point(col_x[0], line_y),
                pymupdf.Point(col_x[-1], line_y),
                width=0.7,
            )
        for x in col_x:
            page.draw_line(
                pymupdf.Point(x, table_top),
                pymupdf.Point(x, table_top + row_h * len(TABLE)),
                width=0.7,
            )
    document.save(path)
    document.close()


def build_docx(path: Path, repeat: int) -> None:
    from docx import Document

    document = Document()
    document.add_heading(TITLE, level=1)
    for _ in range(repeat):
        for heading, lines in SECTIONS:
            document.add_heading(heading, level=2)
            for line in lines:
                document.add_paragraph(line, style="List Bullet")
        document.add_heading("4. Bang thong so ky thuat", level=2)
        table = document.add_table(rows=len(TABLE), cols=len(TABLE[0]))
        table.style = "Table Grid"
        for row_index, row in enumerate(TABLE):
            for col_index, cell in enumerate(row):
                table.cell(row_index, col_index).text = cell
    document.save(path)


if __name__ == "__main__":
    repeat = int(sys.argv[2]) if len(sys.argv) > 2 else 1
    pdf_path = OUT / f"spec_{repeat}x.pdf"
    docx_path = OUT / f"spec_{repeat}x.docx"
    build_pdf(pdf_path, repeat)
    build_docx(docx_path, repeat)
    print(f"{pdf_path}  {pdf_path.stat().st_size:,} bytes")
    print(f"{docx_path}  {docx_path.stat().st_size:,} bytes")
